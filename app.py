import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------- Page setup ----------------------------- #
st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_CHARS = 20000  # safety limit on resume text sent to the model


# ----------------------------- Helpers ----------------------------- #
def get_api_key(sidebar_key: str) -> str:
    """Priority: Streamlit secrets -> environment variable -> sidebar input."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.getenv("GEMINI_API_KEY", "") or sidebar_key


def extract_text(uploaded_file) -> str:
    """Extract plain text from PDF, DOCX or TXT."""
    name = uploaded_file.name.lower()
    data = uploaded_file.read()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description.strip()}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: Not provided. Evaluate the resume for general ATS-friendliness "
        "and overall quality.\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Analyze the resume below{" against the job description" if job_description.strip() else ""} and return ONLY a valid JSON object
with exactly this structure (no markdown, no extra text):

{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "section_scores": {{
    "formatting_and_structure": <0-100>,
    "keywords_and_skills": <0-100>,
    "work_experience_impact": <0-100>,
    "education_and_certifications": <0-100>,
    "readability_and_clarity": <0-100>
  }},
  "keywords_found": ["..."],
  "keywords_missing": ["..."],
  "strengths": ["..."],
  "improvements": [
    {{
      "priority": "High" | "Medium" | "Low",
      "area": "<short area name>",
      "issue": "<what is wrong>",
      "suggestion": "<specific, actionable fix>"
    }}
  ],
  "bullet_rewrites": [
    {{
      "original": "<a weak bullet from the resume>",
      "improved": "<stronger, quantified, action-verb version>"
    }}
  ]
}}

Scoring guidance: be realistic and strict. 90+ is exceptional, 70-89 is good, 50-69 needs work,
below 50 is poor. Consider: contact info, standard section headings, keyword relevance,
quantified achievements, action verbs, length, consistency, and parsing-friendliness.
Give 5-8 improvements ordered by priority and 3 bullet rewrites. Only use facts present in the resume;
never invent experience.

{jd_block}
RESUME:
{resume_text}
"""


def parse_json(text: str) -> dict:
    text = text.strip()
    text = re.sub(r"^```(?:json)?|```$", "", text, flags=re.MULTILINE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def analyze_resume(api_key: str, model: str, resume_text: str, jd: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text[:MAX_CHARS], jd),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_json(response.text)


def score_color(score: int) -> str:
    if score >= 80:
        return "🟢"
    if score >= 60:
        return "🟡"
    return "🔴"


# ----------------------------- Sidebar ----------------------------- #
with st.sidebar:
    st.header("⚙️ Settings")
    sidebar_key = st.text_input(
        "Gemini API Key",
        type="password",
        help="Get a free key at https://aistudio.google.com/apikey. "
        "Not needed if it is set in Streamlit secrets.",
    )
    model_name = st.text_input("Model", value=DEFAULT_MODEL)
    st.caption("Your resume is sent to the Gemini API for analysis and is not stored by this app.")

# ----------------------------- Main UI ----------------------------- #
st.title("📄 ATS Resume Checker")
st.write("Upload your resume to get an ATS score and specific suggestions to improve it.")

col_left, col_right = st.columns(2)
with col_left:
    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
with col_right:
    job_description = st.text_area(
        "Job description (optional, for a tailored score)",
        height=180,
        placeholder="Paste the job description here to see keyword match against the role...",
    )

if st.button("Analyze Resume", type="primary", use_container_width=True):
    api_key = get_api_key(sidebar_key)

    if not api_key:
        st.error("Please provide a Gemini API key in the sidebar or in Streamlit secrets.")
        st.stop()
    if uploaded is None:
        st.warning("Please upload a resume first.")
        st.stop()

    try:
        resume_text = extract_text(uploaded)
    except Exception as e:
        st.error(f"Could not read the file: {e}")
        st.stop()

    if len(resume_text.strip()) < 100:
        st.error(
            "Very little text could be extracted. If your resume is a scanned image or "
            "heavily designed PDF, an ATS would likely struggle to read it too. "
            "Try exporting a text-based PDF or DOCX."
        )
        st.stop()

    with st.spinner("Analyzing your resume with Gemini..."):
        try:
            result = analyze_resume(api_key, model_name, resume_text, job_description)
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            st.stop()

    # ---------------- Results ---------------- #
    score = int(result.get("ats_score", 0))
    st.divider()
    c1, c2 = st.columns([1, 2])
    with c1:
        st.metric("ATS Score", f"{score}/100")
        st.progress(min(max(score, 0), 100) / 100)
        st.markdown(f"## {score_color(score)}")
    with c2:
        st.subheader("Summary")
        st.write(result.get("summary", ""))

    st.subheader("Score Breakdown")
    sections = result.get("section_scores", {})
    cols = st.columns(len(sections) or 1)
    for col, (name, val) in zip(cols, sections.items()):
        with col:
            st.metric(name.replace("_", " ").title(), f"{val}")
            st.progress(min(max(int(val), 0), 100) / 100)

    k1, k2 = st.columns(2)
    with k1:
        st.subheader("✅ Keywords Found")
        found = result.get("keywords_found", [])
        st.write(", ".join(f"`{k}`" for k in found) if found else "None detected.")
    with k2:
        st.subheader("❌ Keywords Missing")
        missing = result.get("keywords_missing", [])
        st.write(", ".join(f"`{k}`" for k in missing) if missing else "None identified.")

    st.subheader("💪 Strengths")
    for s in result.get("strengths", []):
        st.markdown(f"- {s}")

    st.subheader("🛠️ Recommended Improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟡"}
    for item in result.get("improvements", []):
        p = item.get("priority", "Medium")
        with st.expander(f"{icons.get(p, '🟠')} [{p}] {item.get('area', 'General')}"):
            st.markdown(f"**Issue:** {item.get('issue', '')}")
            st.markdown(f"**Fix:** {item.get('suggestion', '')}")

    rewrites = result.get("bullet_rewrites", [])
    if rewrites:
        st.subheader("✍️ Sample Bullet Rewrites")
        for r in rewrites:
            st.markdown(f"**Before:** {r.get('original', '')}")
            st.markdown(f"**After:** {r.get('improved', '')}")
            st.write("")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )
